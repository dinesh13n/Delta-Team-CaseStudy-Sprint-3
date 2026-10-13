"""Build a styled Word briefing from the Markdown source (run from the repository root).
Usage: python deliverables/_build/build_briefing_docx.py deliverables/01-CTO-Briefing-v2.0.md
pandoc converts the Markdown; python-docx then applies the Team Delta styles (fonts, heading colours,
table header shading and borders) so the document reads well without manual formatting."""

import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PETROL, TEAL, GREY, INK = RGBColor(0x0B, 0x3C, 0x49), RGBColor(0x1F, 0x7A, 0x8C), RGBColor(0x5F, 0x6E, 0x71), RGBColor(0x1B, 0x2A, 0x2F)
src = Path(sys.argv[1]).resolve()
out = src.with_suffix(".docx")
DELIV = src.parent


def reference_doc(path: Path) -> None:
    ref = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], capture_output=True, check=True).stdout
    path.write_bytes(ref)
    d = Document(path)
    st = {s.name: s for s in d.styles}  # pandoc's names are case-sensitive; d.styles[...] lower-cases built-ins
    st["Normal"].font.name, st["Normal"].font.size, st["Normal"].font.color.rgb = "Calibri", Pt(10.5), INK
    for name, size, colour in (("Title", 28, PETROL), ("Subtitle", 15, TEAL), ("Heading 1", 18, PETROL), ("Heading 2", 13.5, TEAL), ("Heading 3", 12, TEAL)):
        f = st[name].font
        f.name, f.size, f.color.rgb, f.bold = "Calibri", Pt(size), colour, name != "Subtitle"
    for name in ("Author", "Date"):
        if name in st:
            st[name].font.name, st[name].font.color.rgb = "Calibri", GREY
    for name in ("Body Text", "First Paragraph", "Compact"):
        if name in st:
            st[name].font.name = "Calibri"
    d.save(path)


def shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def borders(table) -> None:
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "C5D6DA")
        b.append(e)
    tblPr.append(b)


with tempfile.TemporaryDirectory() as tmp:
    ref = Path(tmp) / "reference.docx"
    reference_doc(ref)
    subprocess.run(["pandoc", str(src), "-o", str(out), f"--reference-doc={ref}", f"--resource-path={DELIV}"], check=True)

d = Document(out)
for sec in d.sections:
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    footer = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    footer.text = "Team Delta  ·  From Brownfield to Governed AI  ·  CTO Briefing " + (src.stem.split("-")[-1] if "-v" in src.stem else "")
    for r in footer.runs:
        r.font.size, r.font.color.rgb, r.font.name = Pt(8.5), GREY, "Calibri"
for t in d.tables:
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for i, row in enumerate(t.rows):
        for cell in row.cells:
            if i == 0:
                shade(cell, "0B3C49")
            elif i % 2 == 0:
                shade(cell, "F3F7F8")
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    r.font.name = "Calibri"
                    if i == 0:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
for shape in d.inline_shapes:  # fit diagrams to the text width
    ratio = shape.height / shape.width
    shape.width = Cm(17.0)
    shape.height = int(shape.width * ratio)
d.save(out)
print("wrote", out)
