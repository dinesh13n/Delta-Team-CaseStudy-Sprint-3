"""Make 02-CTO-Presentation-v2.2.pptx from the presenter's edited v2.1 deck: insert a problem-statement slide
after the "01 The starting point" divider, and move the version labels to v2.2. Everything else in the
input deck (including the presenter's own edits) is kept as it is.
Usage: python add_problem_slide_v2_2.py INPUT_v2.1.pptx OUTPUT_v2.2.pptx"""

import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

SRC, OUT = sys.argv[1], sys.argv[2]
PETROL, TEAL, RED, GREEN, INK, GREY = "0B3C49", "1F7A8C", "C0392B", "2E8B57", "1B2A2F", "5F6E71"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

prs = Presentation(SRC)
layout = next(l for l in prs.slide_layouts if l.name == "CONTENT")
divider_index = next(i for i, s in enumerate(prs.slides) if s.slide_layout.name == "DARK" and any(sh.has_text_frame and sh.text_frame.text == "The starting point" for sh in s.shapes))
reference_slide = prs.slides[divider_index + 2]  # "62 findings": a CONTENT slide whose slide-number placeholder we copy


def rgb(h):
    return RGBColor.from_string(h)


def text(slide, x, y, w, h, parts, size=14, color=INK, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, spacing=None):
    """parts: list of paragraphs; each paragraph is a list of (text, bold, colour or None)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i:
            p.space_before = Pt(4)
        for t, bold, col in para:
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, "Calibri"
            r.font.color.rgb = rgb(col or color)
            if spacing:
                r.font._element.set("spc", str(spacing))
    return tb


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


slide = prs.slides.add_slide(layout)
# kicker and title, matching the other content slides
text(slide, 0.5, 0.12, 9.0, 0.28, [[("01  THE STARTING POINT", True, TEAL)]], size=11, spacing=200)
title_ph = next(ph for ph in slide.placeholders if ph.placeholder_format.type == 1)
title_ph.text = "Problem statement: what we were handed"
tf_t = title_ph.text_frame
tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0

# left column: context, what we were handed, the task
cards = [
    ("Context", "The repository runs logistics operations: shipments, tracking events, vehicles, routes, carrier bookings and an AI assistant that summarises exceptions for dispatchers.", "E8F1F2", PETROL),
    ("What we were handed (Repo 1.0)", "Three engineering generations side by side: legacy batch scripts, a partly modernised FastAPI service and an AI gateway, with nothing controlling who could do what.", "FCEBEA", RED),
    ("The task", "Turn it into a production-grade Repo 2.0 with evidence for every claim, and capture the business knowledge in a semantic layer that survives a change of AI model.", "E5F3EB", GREEN),
]
for i, (head, body, fill, col) in enumerate(cards):
    y = 1.3 + i * 1.85
    box(slide, 0.5, y, 5.55, 1.7, fill)
    text(slide, 0.72, y + 0.16, 5.1, 1.4, [[(head, True, col)], [(body, False, None)]], size=14)

# right column: the six defects
text(slide, 6.4, 1.3, 6.43, 0.4, [[("Six defects that made it untrustworthy", True, PETROL)]], size=18)
defects = [
    ("Unprovable: ", "the documented test command failed (exit 2), so nobody could prove anything worked."),
    ("Open identity: ", "any caller could claim any role by sending a request header."),
    ("Unguarded AI: ", "the AI endpoint answered with no authorisation and reported its own guardrail as not enforced."),
    ("Wrong data: ", "a lookup that missed returned a different record."),
    ("No accountability: ", "the audit trail recorded time and action only, so no one could reconstruct who did what."),
    ("Empty promises: ", "policy, Terraform and CI files existed but did nothing; the docs promised AI features that did not exist."),
]
for i, (lead, rest) in enumerate(defects):
    y = 1.85 + i * 0.84
    c = box(slide, 6.4, y + 0.06, 0.48, 0.48, RED, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(i + 1)
    r.font.size, r.font.bold, r.font.name = Pt(16), True, "Calibri"
    r.font.color.rgb = rgb("FFFFFF")
    text(slide, 7.05, y, 5.78, 0.6, [[(lead, True, RED), (rest, False, None)]], size=14, anchor=MSO_ANCHOR.MIDDLE)

# slide number placeholder, copied from a neighbouring content slide so numbering stays consistent
for sh in reference_slide.shapes:
    if sh.is_placeholder and "Slide Number" in sh.name:
        slide.shapes._spTree.append(copy.deepcopy(sh._element))

notes = slide.notes_slide
if notes.notes_text_frame is None:  # pptxgenjs notes master has no body placeholder: copy the one a sibling slide uses
    src_body = next(sh for sh in reference_slide.notes_slide.shapes if sh.is_placeholder and sh.placeholder_format.type == 2)
    notes.shapes._spTree.append(copy.deepcopy(src_body._element))
notes.notes_text_frame.text = (
    "Set the scene before any architecture. Context: a logistics operations repository with an AI assistant for dispatchers. "
    "What we were handed: three code generations and nothing controlling who could do what. Walk the six defects quickly; each one "
    "maps to findings we later fixed and proved. Then the task: production-grade Repo 2.0 with evidence for every claim, and a "
    "semantic layer that survives a change of AI model. Next slide shows the same picture as an architecture diagram."
)

# move the new slide to sit right after the divider, and put it in the same section
sldIdLst = prs.slides._sldIdLst
new_id = sldIdLst[-1]
sldIdLst.remove(new_id)
sldIdLst.insert(divider_index + 1, new_id)
divider_id = sldIdLst[divider_index].get("id")
for sec_list in prs.part._element.iter(f"{{{P14}}}sldIdLst"):
    ids = [e.get("id") for e in sec_list]
    if divider_id in ids:
        entry = etree.SubElement(sec_list, f"{{{P14}}}sldId")
        entry.set("id", new_id.get("id"))
        sec_list.remove(entry)
        sec_list.insert(ids.index(divider_id) + 1, entry)

# version labels: v2.1 -> v2.2 on the layout footer, the title slide, and the agenda note
for lay in prs.slide_layouts:
    for sh in lay.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if "v2.1" in r.text:
                        r.text = r.text.replace("v2.1", "v2.2")
for s in prs.slides:
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if "·  v2.1" in r.text:
                    r.text = r.text.replace("·  v2.1", "·  v2.2")
                if r.text.startswith("New in v2.0: part 03"):
                    r.text = r.text + ". New in v2.2: the problem statement (slide 6)"
prs.core_properties.title = "From Brownfield to Governed AI: CTO briefing v2.2"
prs.save(OUT)
print("wrote", OUT, "slides", len(prs.slides))
